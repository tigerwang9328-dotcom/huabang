"""
退货数据统计 API  /api/v1/finance/return-stats/*
- POST /compute      ：按下单时间窗口异步计算各店退货数据（长任务，submit_async_task + 轮询）
- POST /apply-params ：把各店"预估退款率"一键写入日报参数 estimated_return_rate（选中月份）
- POST /export       ：把前端已算好的结果导出为 xlsx
"""
import io
import urllib.parse
from datetime import date, timedelta
from fastapi import APIRouter, Depends, HTTPException, BackgroundTasks, Query
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text

import asyncio
from app.api.v1.deps import get_db, get_current_user
from app.models.sys import SysUser as User
from app.api.v1.tasks import submit_async_task
from app.services.finance.return_stats_util import (
    compute_return_stats_sync, compute_spu_return_stats_sync,
    read_spu_return_cache, compute_and_cache_spu_return_default,
    read_store_return_cache, compute_and_cache_store_return_default,
)

router = APIRouter(prefix="/finance/return-stats")


@router.post("/compute", status_code=202)
async def compute(
    body: dict,
    background: BackgroundTasks,
    _ = Depends(get_current_user),
):
    """异步计算退货数据。body: {start_date, end_date}（含两端，按下单时间）。返回 task_id，轮询 /tasks/status/{id} 取 result。"""
    sd = body.get("start_date")
    ed = body.get("end_date")
    if not sd or not ed:
        raise HTTPException(status_code=422, detail="start_date 和 end_date 必填(YYYY-MM-DD)")
    try:
        d0 = date.fromisoformat(str(sd))
        d1 = date.fromisoformat(str(ed))
    except (ValueError, TypeError):
        raise HTTPException(status_code=422, detail="日期格式错误，须为 YYYY-MM-DD")
    if d1 < d0:
        raise HTTPException(status_code=422, detail="结束日期不能早于开始日期")
    end_excl = (d1 + timedelta(days=1)).isoformat()
    return submit_async_task(
        background, "return_stats",
        lambda: asyncio.to_thread(compute_return_stats_sync, d0.isoformat(), end_excl),
        start_date=d0.isoformat(), end_date=d1.isoformat(),
    )


@router.get("/spu-cache")
async def spu_return_cache(_ = Depends(get_current_user)):
    """读取每日定时预算的「款式退货退款」默认窗口(前20~10天)缓存，秒开。
    页面常态展示用；缓存不存在返回 {cached: false}。"""
    data = read_spu_return_cache()
    if data is None:
        return {"cached": False}
    return data


@router.post("/spu-cache/refresh")
async def spu_return_cache_refresh(_ = Depends(get_current_user)):
    """手动立即重算默认窗口缓存(约数秒~十几秒)，并返回最新结果。"""
    return await asyncio.to_thread(compute_and_cache_spu_return_default)


@router.get("/store-cache")
async def store_return_cache(_ = Depends(get_current_user)):
    """读取每日定时预算的「店铺退货数据」默认窗口(前20~10天)缓存，秒开。
    页面常态展示用；缓存不存在返回 {cached: false}。"""
    data = read_store_return_cache()
    if data is None:
        return {"cached": False}
    return data


@router.post("/store-cache/refresh")
async def store_return_cache_refresh(_ = Depends(get_current_user)):
    """手动立即重算默认窗口店铺级缓存(约数秒~十几秒)，并返回最新结果。"""
    return await asyncio.to_thread(compute_and_cache_store_return_default)


@router.get("/spu")
async def spu_return_stats(
    date_from: str = Query(..., description="下单时间起(含) YYYY-MM-DD"),
    date_to: str = Query(..., description="下单时间止(含) YYYY-MM-DD"),
    _ = Depends(get_current_user),
):
    """款式(SPU)退货退款查询。与店铺版同口径(下单时间口径)，维度改为款式、按订单数计数、整单归属。
    耗时约数秒(DWD 聚合 + 退款类型解析)，直接同步返回 result，无需轮询。"""
    try:
        d0 = date.fromisoformat(str(date_from))
        d1 = date.fromisoformat(str(date_to))
    except (ValueError, TypeError):
        raise HTTPException(status_code=422, detail="日期格式错误，须为 YYYY-MM-DD")
    if d1 < d0:
        raise HTTPException(status_code=422, detail="结束日期不能早于开始日期")
    end_excl = (d1 + timedelta(days=1)).isoformat()
    return await asyncio.to_thread(compute_spu_return_stats_sync, d0.isoformat(), end_excl)


@router.post("/apply-params")
async def apply_params(
    body: dict,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """把各店预估退款率写入 finance_store_daily_params.estimated_return_rate（按月 UPSERT）。"""
    month = body.get("month")
    rows = body.get("rows") or []
    if not month or not rows:
        raise HTTPException(status_code=422, detail="month 和 rows 必填")
    count = 0
    for r in rows:
        sid = r.get("store_id")
        rate = r.get("est_refund_rate")
        if not sid or rate is None:
            continue
        await db.execute(text("""
            INSERT INTO finance_store_daily_params (month, store_id, store_name, estimated_return_rate, updated_by)
            VALUES (:m, :sid, :nm, :rate, :uid)
            ON CONFLICT (month, store_id) DO UPDATE SET
                estimated_return_rate = :rate,
                updated_by = :uid,
                updated_at = NOW()
        """), {"m": month, "sid": int(sid), "nm": r.get("store_name"),
               "rate": float(rate), "uid": current_user.id})
        count += 1
    await db.commit()
    return {"ok": True, "month": month, "count": count}


@router.post("/export")
async def export(body: dict, _ = Depends(get_current_user)):
    """把前端已算好的退货数据结果导出为 xlsx。body 即 /compute 的 result。"""
    rows = body.get("rows") or []
    totals = body.get("totals") or {}
    sd = body.get("start_date") or ""
    ed = body.get("end_date") or body.get("end_date_excl") or ""
    ca = body.get("computed_at") or ""
    ao = body.get("data_as_of") or ""

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "退货数据"
    hdr = ["店铺组", "店铺", "已发货待退款", "已发货已退款", "已发货订单数", "订单总数", "预估退款率", "仅退款率"]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(hdr))
    t = ws.cell(row=1, column=1, value=f"退货数据（下单时间口径）  {sd} ~ {ed}    数据截至：{ao}    计算时间：{ca}")
    t.font = Font(bold=True, size=12)
    t.alignment = Alignment(horizontal="left", vertical="center")

    ws.append(hdr)
    fill = PatternFill("solid", fgColor="1F4E79")
    group_fill = PatternFill("solid", fgColor="E8F0FE")
    total_fill = PatternFill("solid", fgColor="E8F0FE")
    hf = Font(bold=True, color="FFFFFF")
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=2, column=c)
        cell.fill = fill; cell.font = hf; cell.alignment = Alignment(horizontal="center")

    r = 3
    for row in rows:
        is_group = row.get("_row_kind") == "group"
        level = int(row.get("_level") or 0)
        group_name = str(row.get("_group_name") or "")
        store_name = str(row.get("store_name") or "")
        if is_group:
            store_name = "合计"
        elif level > 0:
            store_name = ("　" * level) + store_name
        ws.cell(row=r, column=1, value=group_name)
        ws.cell(row=r, column=2, value=store_name)
        ws.cell(row=r, column=3, value=row.get("pending_refund"))
        ws.cell(row=r, column=4, value=row.get("refunded"))
        ws.cell(row=r, column=5, value=row.get("shipped"))
        ws.cell(row=r, column=6, value=row.get("total"))
        ws.cell(row=r, column=7, value=round(float(row.get("est_refund_rate") or 0), 6)).number_format = "0.00%"
        ws.cell(row=r, column=8, value=round(float(row.get("only_refund_rate") or 0), 6)).number_format = "0.00%"
        if is_group:
            for c in range(1, len(hdr) + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = group_fill
                cell.font = Font(bold=True)
        r += 1

    # 合计
    ws.cell(row=r, column=2, value="合计")
    ws.cell(row=r, column=3, value=totals.get("pending_refund"))
    ws.cell(row=r, column=4, value=totals.get("refunded"))
    ws.cell(row=r, column=5, value=totals.get("shipped"))
    ws.cell(row=r, column=6, value=totals.get("total"))
    tot = totals.get("total") or 0
    if tot:
        ws.cell(row=r, column=7, value=round((float(totals.get("pending_refund") or 0) + float(totals.get("refunded") or 0)) / tot, 6)).number_format = "0.00%"
        ws.cell(row=r, column=8, value=round((tot - float(totals.get("shipped") or 0)) / tot, 6)).number_format = "0.00%"
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = Font(bold=True)
        cell.fill = total_fill

    for i, w in enumerate([18, 42, 14, 14, 14, 10, 11, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    fn = urllib.parse.quote(f"退货数据_{sd}_{ed}.xlsx")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"return_stats.xlsx\"; filename*=UTF-8''{fn}"},
    )
