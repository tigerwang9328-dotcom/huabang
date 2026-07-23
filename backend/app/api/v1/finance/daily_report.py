"""
日报 API - /api/v1/finance/daily-report/*
"""
from datetime import date
import urllib.parse
from typing import Optional
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text

from app.api.v1.deps import get_db, get_current_user
from app.models.sys import SysUser as User
from app.services.finance.daily_report_service import get_daily_report
from app.core.permissions import check_permission

router = APIRouter(prefix="/finance/daily-report")


@router.get("")
async def daily_report(
    biz_date: date = Query(default=None, description="YYYY-MM-DD，默认昨天"),
    store_id: Optional[int] = Query(default=None),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """获取指定日期完整日报（全店或单店）"""
    from datetime import timedelta
    from zoneinfo import ZoneInfo
    from datetime import datetime
    if not biz_date:
        biz_date = datetime.now(ZoneInfo("Asia/Shanghai")).date() - timedelta(days=1)
    return await get_daily_report(db, biz_date, store_id)


@router.get("/export")
async def export_daily_report(
    biz_date: date = Query(default=None, description="YYYY-MM-DD"),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """下载日报 Excel"""
    if not check_permission(current_user, "finance:export"):
        raise HTTPException(status_code=403, detail="无权限: finance:export")
    from datetime import timedelta
    from zoneinfo import ZoneInfo
    from datetime import datetime
    if not biz_date:
        biz_date = datetime.now(ZoneInfo("Asia/Shanghai")).date() - timedelta(days=1)

    report = await get_daily_report(db, biz_date)
    content = _build_excel(report)
    fn_ascii = f"daily_report_{biz_date}.xlsx"
    fn_utf8 = urllib.parse.quote(f"\u7267\u9a6c\u4eba\u65e5\u62a5_{biz_date}.xlsx")
    disp = f'attachment; filename="{fn_ascii}"; filename*=UTF-8\'\'{fn_utf8}'
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": disp},
    )


# ─── 参数维护 CRUD ─────────────────────────────────────────────────

@router.get("/params")
async def list_params(
    month: str = Query(description="YYYY-MM"),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """获取指定月份所有店铺参数"""
    r = await db.execute(text("""
        SELECT
            p.id, p.month, p.store_id, p.store_name, p.platform,
            p.platform_commission_rate, p.platform_income_rate,
            p.estimated_return_rate,
            p.refund_only_rate,
            p.freight_insurance_unit_cost, p.express_unit_cost,
            p.package_unit_cost, p.return_labor_unit_cost,
            p.goods_loss_unit_cost, p.promotion_unit_cost,
            p.return_rate_warning_threshold, p.warning_enabled,
            p.remark, p.updated_at
        FROM finance_store_daily_params p
        WHERE p.month = :month
        ORDER BY p.store_name
    """), {"month": month})
    rows = r.mappings().all()
    return [dict(row) for row in rows]


@router.get("/params/stores")
async def list_stores_for_params(
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """返回所有活跃店铺（用于参数维护页初始化）"""
    r = await db.execute(text("""
        SELECT s.id AS store_id, s.store_name, COALESCE(bp.name,'') AS platform
        FROM biz_stores s
        LEFT JOIN biz_platforms bp ON bp.id = s.platform_id
        WHERE s.is_active = true
        ORDER BY s.store_name
    """))
    rows = r.mappings().all()
    return [dict(row) for row in rows]


@router.post("/params")
async def upsert_param(
    body: dict,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """新增或更新单店参数（按 month+store_id UPSERT）"""
    month = body.get("month")
    store_id = body.get("store_id")
    if not month or not store_id:
        raise HTTPException(status_code=422, detail="month 和 store_id 必填")

    fields = [
        "platform_commission_rate", "platform_income_rate", "estimated_return_rate",
        "refund_only_rate",
        "freight_insurance_unit_cost", "express_unit_cost", "package_unit_cost",
        "return_labor_unit_cost", "goods_loss_unit_cost", "promotion_unit_cost",
        "return_rate_warning_threshold", "warning_enabled", "remark",
        "store_name", "platform",
    ]
    set_parts = ", ".join(f"{f} = :{f}" for f in fields if f in body)
    params = {f: body[f] for f in fields if f in body}
    params["month"] = month
    params["store_id"] = store_id
    params["updated_by"] = current_user.id

    await db.execute(text(f"""
        INSERT INTO finance_store_daily_params (month, store_id, {', '.join(f for f in fields if f in body)}, updated_by)
        VALUES (:month, :store_id, {', '.join(':' + f for f in fields if f in body)}, :updated_by)
        ON CONFLICT (month, store_id) DO UPDATE SET
            {set_parts},
            updated_by = :updated_by,
            updated_at = NOW()
    """), params)
    await db.commit()
    return {"ok": True}


@router.post("/params/batch")
async def batch_upsert_params(
    body: dict,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """批量 UPSERT 参数（body: {month, rows: [{store_id, ...}]}）"""
    month = body.get("month")
    rows = body.get("rows", [])
    if not month or not rows:
        raise HTTPException(status_code=422, detail="month 和 rows 必填")
    for row in rows:
        row["month"] = month
        row["updated_by"] = current_user.id
        await upsert_param(row, db, current_user)
    return {"ok": True, "count": len(rows)}


@router.post("/params/copy-month")
async def copy_month_params(
    body: dict,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """将上月参数复制到目标月。

    目标月不存在的店铺直接新增；目标月已有但仍是默认空参数的行会被填充。
    已经手工配置过的目标月参数不覆盖。
    """
    src = body.get("src_month")
    dst = body.get("dst_month")
    if not src or not dst:
        raise HTTPException(status_code=422, detail="src_month 和 dst_month 必填")

    params = {"src": src, "dst": dst, "uid": current_user.id}
    r = await db.execute(text("""
        WITH src AS (
            SELECT *
            FROM finance_store_daily_params
            WHERE month = :src
        ),
        inserted AS (
            INSERT INTO finance_store_daily_params (
                month, store_id, store_name, platform,
                platform_commission_rate, platform_income_rate, estimated_return_rate,
                refund_only_rate,
                freight_insurance_unit_cost, express_unit_cost, package_unit_cost,
                return_labor_unit_cost, goods_loss_unit_cost, promotion_unit_cost,
                return_rate_warning_threshold, warning_enabled, remark, updated_by
            )
            SELECT
                :dst, store_id, store_name, platform,
                platform_commission_rate, platform_income_rate, estimated_return_rate,
                refund_only_rate,
                freight_insurance_unit_cost, express_unit_cost, package_unit_cost,
                return_labor_unit_cost, goods_loss_unit_cost, promotion_unit_cost,
                return_rate_warning_threshold, warning_enabled, remark, :uid
            FROM src
            WHERE true
            ON CONFLICT (month, store_id) DO NOTHING
            RETURNING id
        ),
        updated AS (
            UPDATE finance_store_daily_params dst
            SET
                store_name = COALESCE(NULLIF(dst.store_name, ''), src.store_name),
                platform = COALESCE(NULLIF(dst.platform, ''), src.platform),
                platform_commission_rate = src.platform_commission_rate,
                platform_income_rate = src.platform_income_rate,
                estimated_return_rate = src.estimated_return_rate,
                refund_only_rate = src.refund_only_rate,
                freight_insurance_unit_cost = src.freight_insurance_unit_cost,
                express_unit_cost = src.express_unit_cost,
                package_unit_cost = src.package_unit_cost,
                return_labor_unit_cost = src.return_labor_unit_cost,
                goods_loss_unit_cost = src.goods_loss_unit_cost,
                promotion_unit_cost = src.promotion_unit_cost,
                return_rate_warning_threshold = src.return_rate_warning_threshold,
                warning_enabled = src.warning_enabled,
                remark = src.remark,
                updated_by = :uid,
                updated_at = NOW()
            FROM src
            WHERE dst.month = :dst
              AND dst.store_id = src.store_id
              AND NOT EXISTS (
                  SELECT 1
                  FROM inserted i
                  WHERE i.id = dst.id
              )
              AND COALESCE(dst.platform_commission_rate, 0) = 0
              AND COALESCE(dst.platform_income_rate, 0) IN (0, 1)
              AND COALESCE(dst.estimated_return_rate, 0) = 0
              AND COALESCE(dst.refund_only_rate, 0) = 0
              AND COALESCE(dst.freight_insurance_unit_cost, 0) = 0
              AND COALESCE(dst.express_unit_cost, 0) = 0
              AND COALESCE(dst.package_unit_cost, 0) = 0
              AND COALESCE(dst.return_labor_unit_cost, 0) = 0
              AND COALESCE(dst.goods_loss_unit_cost, 0) = 0
              AND COALESCE(dst.promotion_unit_cost, 0) = 0
              AND COALESCE(dst.return_rate_warning_threshold, 0) IN (0, 0.08)
              AND COALESCE(dst.warning_enabled, true) = true
              AND COALESCE(dst.remark, '') = ''
            RETURNING dst.id
        )
        SELECT
            (SELECT count(*) FROM inserted) AS inserted,
            (SELECT count(*) FROM updated) AS updated
    """), params)
    await db.commit()
    result = r.mappings().one()
    inserted = result["inserted"]
    updated = result["updated"]
    return {"ok": True, "copied": inserted + updated, "inserted": inserted, "updated": updated}


# ─── Excel 构建 ─────────────────────────────────────────────────────

def _build_excel(report: dict) -> bytes:
    """构建日报 Excel，格式对齐参考 xlsx"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import io
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl 未安装")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"日报_{report['biz_date']}"

    # ── 样式定义 ──
    header_fill = PatternFill("solid", fgColor="1F4E79")
    sub_fill    = PatternFill("solid", fgColor="2F6FED")
    sum_fill    = PatternFill("solid", fgColor="E8F0FE")
    hdr_font    = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
    sub_font    = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
    sum_font    = Font(name="微软雅黑", bold=True, size=10)
    data_font   = Font(name="微软雅黑", size=10)
    thin        = Side(style="thin", color="D1D5DB")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_al   = Alignment(horizontal="center", vertical="center")
    right_al    = Alignment(horizontal="right",  vertical="center")
    left_al     = Alignment(horizontal="left",   vertical="center")

    # ── 表头（第2行）──
    headers = [
        ("序号",                  8, center_al),
        ("店铺组",               18, left_al),
        ("店铺",                 28, left_al),
        ("利润",                 12, right_al),
        ("费用合计",             12, right_al),
        ("广告费",               11, right_al),
        ("毛利",                  9, right_al),
        ("每单利润",             10, right_al),
        ("发货订单金额",         13, right_al),
        ("预估退款率",           10, right_al),
        ("退款后金额",           13, right_al),
        ("实际金额\n（平台扣点）", 14, right_al),
        ("成交\n订单数",          10, right_al),
        ("产品\n件数",            10, right_al),
        ("产品成本",             12, right_al),
        ("退货\n产品成本",        12, right_al),
        ("运费险",               10, right_al),
        ("包装成本",             10, right_al),
        ("快递费",               10, right_al),
        ("发货赔付、其他赔付",   18, right_al),
        ("推广成本",             10, right_al),
        ("退货人工",             10, right_al),
        ("货值损耗(每单1元)",    15, right_al),
        ("仅退款率",             10, right_al),
        ("实际\n订单数",          10, right_al),
        ("产品\n件数",            10, right_al),
        ("客单价",               10, right_al),
        ("每单\n广告费",          10, right_al),
        ("roi",                   9, right_al),
    ]

    # ── 标题行 ──
    last_col = get_column_letter(len(headers))
    ws.merge_cells(f"A1:{last_col}1")
    title_cell = ws["A1"]
    title_cell.value = f"牧马人服饰 · {report['biz_date']} 经营日报"
    title_cell.font  = Font(name="微软雅黑", bold=True, size=14, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    for col_idx, (label, width, _al) in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=label)
        cell.font = sub_font
        cell.fill = sub_fill
        cell.alignment = center_al
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 20

    # ── 数据行（有序行：成员店行 + 组合计行，与页面同序）──
    rows = report.get("report_rows") or report.get("stores", [])
    store_seq = 0
    for r_idx, row in enumerate(rows, 3):
        is_total = row.get("row_kind") == "group_total"
        if not is_total:
            store_seq += 1
        cur_font = sum_font if is_total else data_font
        row_fill = sum_fill if is_total else None
        level = int(row.get("level", 0) or 0)

        def wcell(col, val, fmt=None, al=right_al):
            c = ws.cell(row=r_idx, column=col, value=val)
            c.font = cur_font
            c.alignment = al
            c.border = border
            if fmt:
                c.number_format = fmt
            if row_fill:
                c.fill = row_fill
            return c

        income_rate = float(row.get("platform_income_rate") or 1.0)

        group_name = str(row.get("group_name") or "")
        group_label = ("　" * max(level - (0 if is_total else 1), 0)) + group_name if group_name else ""
        store_label = "合计" if is_total else ("　" * level) + str(row["store_name"])

        wcell(1,  "合计" if is_total else store_seq, al=center_al)
        wcell(2,  group_label, al=left_al)
        wcell(3,  store_label, al=left_al)
        wcell(4,  f"=L{r_idx}-O{r_idx}+P{r_idx}-Q{r_idx}-R{r_idx}-S{r_idx}-T{r_idx}-U{r_idx}-V{r_idx}-W{r_idx}-F{r_idx}", "#,##0.00")
        wcell(5,  f"=Q{r_idx}+R{r_idx}+S{r_idx}+T{r_idx}+U{r_idx}+V{r_idx}+F{r_idx}+W{r_idx}", "#,##0.00")
        wcell(6,  row["ad_cost"],              "#,##0")
        wcell(7,  f"=IFERROR(D{r_idx}/I{r_idx},0)", "0.00%")
        wcell(8,  f"=IFERROR(D{r_idx}/M{r_idx},0)", "#,##0.00")
        wcell(9,  row["sale_amount"],          "#,##0.00")
        wcell(10, row["refund_rate"],          "0%")
        wcell(
            11,
            row["refund_net_amount"] if is_total else f"=I{r_idx}*(1-J{r_idx})",
            "#,##0.00",
        )
        wcell(
            12,
            row["platform_net_amount"] if is_total else f"=K{r_idx}*{income_rate:.6f}",
            "#,##0.00",
        )
        wcell(13, row["order_count"])
        wcell(14, row["shipped_qty"])
        wcell(15, row["sale_cogs"],            "#,##0.00")
        wcell(
            16,
            row["refund_cogs_back"] if is_total else f"=O{r_idx}*J{r_idx}",
            "#,##0.00",
        )
        wcell(17, row["freight_insurance"],    "#,##0.00")
        wcell(18, row["package_cost"],         "#,##0.00")
        wcell(19, row["express_cost"],         "#,##0.00")
        wcell(20, row.get("compensation_amount", 0), "#,##0.00")
        wcell(21, row["promotion_cost"],       "#,##0.00")
        wcell(22, row["return_labor_cost"],    "#,##0.00")
        wcell(23, row["goods_loss_cost"],      "#,##0.00")
        wcell(24, row.get("refund_only_rate", None), "0%")
        wcell(25, row.get("actual_order_count", 0))
        wcell(26, row.get("actual_product_qty", 0))
        wcell(27, f"=IFERROR(I{r_idx}/M{r_idx},0)", "#,##0.00")
        wcell(28, f"=IFERROR(F{r_idx}/Z{r_idx},0)", "#,##0.00")
        wcell(29, f"=IFERROR(I{r_idx}/F{r_idx},0)", "0.00")

    # ── 合计行 ──
    s = report.get("summary", {})
    if s:
        sum_row = len(rows) + 3
        ws.row_dimensions[sum_row].height = 18
        def scell(col, val, fmt=None, al=right_al):
            c = ws.cell(row=sum_row, column=col, value=val)
            c.font = sum_font
            c.fill = sum_fill
            c.alignment = al
            c.border = border
            if fmt:
                c.number_format = fmt
            return c
        scell(1, "合计",                   al=center_al)
        scell(2, "",                       al=left_al)
        scell(3, "合计",                   al=left_al)
        scell(4,  s.get("profit"),         "#,##0.00")
        scell(5,  s.get("total_expense"),  "#,##0.00")
        scell(6,  s.get("ad_cost"),        "#,##0")
        scell(7,  f"=IFERROR(D{sum_row}/I{sum_row},0)", "0.00%")
        scell(8,  f"=IFERROR(D{sum_row}/M{sum_row},0)", "#,##0.00")
        scell(9,  s.get("sale_amount"),    "#,##0.00")
        scell(10, s.get("refund_rate"),    "0%")
        scell(11, s.get("refund_net_amount"), "#,##0.00")
        scell(12, s.get("platform_net_amount"), "#,##0.00")
        scell(13, s.get("order_count"))
        scell(14, s.get("shipped_qty"))
        scell(15, s.get("sale_cogs"),      "#,##0.00")
        scell(16, s.get("refund_cogs_back"), "#,##0.00")
        scell(17, s.get("freight_insurance"), "#,##0.00")
        scell(18, s.get("package_cost"),   "#,##0.00")
        scell(19, s.get("express_cost"),   "#,##0.00")
        scell(20, s.get("compensation_amount", 0), "#,##0.00")
        scell(21, s.get("promotion_cost"), "#,##0.00")
        scell(22, s.get("return_labor_cost"), "#,##0.00")
        scell(23, s.get("goods_loss_cost"), "#,##0.00")
        scell(24, s.get("refund_only_rate"), "0%")
        scell(25, s.get("actual_order_count"))
        scell(26, s.get("actual_product_qty"))
        scell(27, f"=IFERROR(I{sum_row}/M{sum_row},0)", "#,##0.00")
        scell(28, f"=IFERROR(F{sum_row}/Z{sum_row},0)", "#,##0.00")
        scell(29, f"=IFERROR(I{sum_row}/F{sum_row},0)", "0.00")

    ws.freeze_panes = "D3"
    ws.print_area = f"A1:{last_col}{ws.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
