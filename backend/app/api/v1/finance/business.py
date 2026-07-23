"""
业务台账 API - 出纳/资产/发票/工资
对标虎狼 routes.py 中 cashier/assets/invoices/payroll 相关路由
"""
from typing import Optional
from datetime import date
from fastapi import APIRouter, Depends, Query, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, text
from pydantic import BaseModel

from app.api.v1.deps import get_db, get_current_user
from app.models.sys import SysUser as User
from app.models.finance.business import (
    FinCashAccount, FinCashFlow, FinFixedAsset, FinInvoice, FinPayrollRecord,
)
from app.models.finance.settings import FinAuditLog

router = APIRouter(prefix="/finance")


# ── 出纳 ──

@router.get("/cashier/accounts")
async def list_cashier_accounts(
    book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    result = await db.execute(
        select(FinCashAccount).where(FinCashAccount.book_id == book_id)
    )
    return {"rows": [
        {"id": a.id, "account_name": a.account_name, "account_type": a.account_type,
         "bank_name": a.bank_name, "account_no": a.account_no,
         "balance": float(a.balance or 0), "is_active": a.is_active}
        for a in result.scalars()
    ]}


class CashAccountCreate(BaseModel):
    account_name: str
    account_type: str = "bank"
    bank_name: Optional[str] = None
    account_no: Optional[str] = None
    balance: float = 0

@router.post("/cashier/accounts", status_code=201)
async def create_cashier_account(
    body: CashAccountCreate, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    a = FinCashAccount(book_id=book_id, **body.model_dump())
    db.add(a)
    await db.flush()
    return {"id": a.id}


@router.get("/cashier/flows")
async def list_cashier_flows(
    book_id: int = Query(1),
    account_id: Optional[int] = Query(None),
    limit: int = Query(50, le=200),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    sql = "SELECT * FROM fin_cash_flows WHERE book_id = :bid"
    params = {"bid": book_id}
    if account_id:
        sql += " AND cash_account_id = :aid"
        params["aid"] = account_id
    sql += " ORDER BY flow_date DESC LIMIT :limit"
    params["limit"] = limit
    result = await db.execute(text(sql), params)
    total_in = total_out = 0
    rows = []
    for r in result.mappings():
        amt = float(r["amount"] or 0)
        if r["flow_type"] == "in":
            total_in += amt
        else:
            total_out += amt
        rows.append(dict(r))
    return {"rows": rows, "total_in": round(total_in, 2), "total_out": round(total_out, 2)}




# ── 出纳 KPI 统计 ──

@router.get("/cashier/stats")
async def cashier_stats(
    book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """出纳管理 KPI：账户总数、总余额、今日/本月收支"""
    from datetime import date as _date
    today = _date.today()
    month_start = _date(today.year, today.month, 1)

    acct_r = await db.execute(text("""
        SELECT COUNT(*) AS cnt, COALESCE(SUM(balance), 0) AS total
        FROM fin_cash_accounts WHERE book_id = :bid AND is_active = true
    """), {"bid": book_id})
    acct = acct_r.mappings().first()

    flow_r = await db.execute(text("""
        SELECT
            COALESCE(SUM(CASE WHEN flow_date = :today AND flow_type='in'  THEN amount END), 0) AS today_in,
            COALESCE(SUM(CASE WHEN flow_date = :today AND flow_type='out' THEN amount END), 0) AS today_out,
            COALESCE(SUM(CASE WHEN flow_date >= :ms   AND flow_type='in'  THEN amount END), 0) AS month_in,
            COALESCE(SUM(CASE WHEN flow_date >= :ms   AND flow_type='out' THEN amount END), 0) AS month_out
        FROM fin_cash_flows WHERE book_id = :bid
    """), {"bid": book_id, "today": today, "ms": month_start})
    flow = flow_r.mappings().first()

    return {
        "account_count":  int(acct["cnt"]),
        "total_balance":  float(acct["total"]),
        "today_in":       float(flow["today_in"]),
        "today_out":      float(flow["today_out"]),
        "month_in":       float(flow["month_in"]),
        "month_out":      float(flow["month_out"]),
    }


class CashFlowCreate(BaseModel):
    cash_account_id: int
    flow_date: str          # YYYY-MM-DD
    flow_type: str          # in / out
    amount: float
    category: Optional[str] = None
    counterpart: Optional[str] = None
    remark: Optional[str] = None


@router.post("/cashier/flows", status_code=201)
async def create_cashier_flow(
    body: CashFlowCreate,
    book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """新增出纳流水并自动更新账户余额"""
    from datetime import date as _date
    flow = FinCashFlow(
        book_id=book_id,
        cash_account_id=body.cash_account_id,
        flow_date=_date.fromisoformat(body.flow_date),
        flow_type=body.flow_type,
        amount=body.amount,
        category=body.category,
        counterpart=body.counterpart,
        remark=body.remark,
    )
    db.add(flow)
    delta = body.amount if body.flow_type == "in" else -body.amount
    await db.execute(
        text("UPDATE fin_cash_accounts SET balance = balance + :d WHERE id = :id AND book_id = :bid"),
        {"d": delta, "id": body.cash_account_id, "bid": book_id},
    )
    await db.flush()
    return {"id": flow.id, "ok": True}


# ── 固定资产 ──

@router.get("/assets")
async def list_assets(
    book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    period = await _asset_current_period(db, book_id)
    result = await db.execute(
        select(FinFixedAsset).where(FinFixedAsset.book_id == book_id)
    )
    rows = []
    for a in result.scalars():
        _sync_asset_depre(a, period)
        rows.append({
            "id": a.id, "asset_code": a.asset_code, "asset_name": a.asset_name,
            "category": a.category, "purchase_date": str(a.purchase_date) if a.purchase_date else "",
            "original_value": float(a.original_value or 0),
            "net_value": float(a.net_value or 0),
            "monthly_depre": float(a.monthly_depre or 0),
            "accumulated_depre": float(a.accumulated_depre or 0),
            "depre_years": a.depre_years,
            "status": a.status,
        })
    return {"rows": rows, "total": len(rows)}


class AssetCreate(BaseModel):
    asset_code: Optional[str] = None
    asset_name: str
    category: Optional[str] = None
    purchase_date: Optional[str] = None
    original_value: float = 0
    depre_method: str = "straight_line"
    depre_years: int = 5
    useful_life_years: Optional[int] = None
    status: Optional[str] = None


def _asset_years(body: AssetCreate) -> int:
    years = body.useful_life_years if body.useful_life_years is not None else body.depre_years
    return max(int(years or 1), 1)


async def _asset_current_period(db: AsyncSession, book_id: int) -> str:
    result = await db.execute(text("SELECT current_period FROM fin_books WHERE id=:bid"), {"bid": book_id})
    period = result.scalar()
    if period:
        return period
    today = date.today()
    return f"{today.year}-{today.month:02d}"


def _elapsed_depre_amount(original_value: float, purchase_date: Optional[date], years: int, period: str) -> float:
    if not purchase_date or not period:
        return 0.0
    try:
        y, m = int(period[:4]), int(period[5:7])
    except Exception:
        today = date.today()
        y, m = today.year, today.month
    elapsed_months = (y * 12 + m) - (purchase_date.year * 12 + purchase_date.month)
    depreciable_months = max(min(elapsed_months, years * 12), 0)
    if depreciable_months <= 0 or years <= 0:
        return 0.0
    monthly = round(float(original_value or 0) / (years * 12), 2)
    return round(min(float(original_value or 0), monthly * depreciable_months), 2)


def _sync_asset_depre(asset: FinFixedAsset, period: str) -> None:
    original = float(asset.original_value or 0)
    years = int(asset.depre_years or 1)
    expected = _elapsed_depre_amount(original, asset.purchase_date, years, period)
    current = float(asset.accumulated_depre or 0)
    accumulated = max(current, expected)
    asset.accumulated_depre = accumulated
    asset.net_value = max(original - accumulated, 0)
    if asset.net_value <= 0 and original > 0:
        asset.status = "scrapped"


async def _next_asset_code(db: AsyncSession, book_id: int) -> str:
    prefix = f"FA-{date.today().strftime('%Y%m%d')}-"
    result = await db.execute(text("""
        SELECT asset_code FROM fin_fixed_assets
        WHERE book_id = :bid AND asset_code LIKE :prefix
        ORDER BY asset_code DESC LIMIT 1
    """), {"bid": book_id, "prefix": prefix + "%"})
    last = result.scalar()
    seq = int(last.rsplit("-", 1)[-1]) + 1 if last else 1
    return f"{prefix}{seq:03d}"


@router.post("/assets", status_code=201)
async def create_asset(
    body: AssetCreate, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    years = _asset_years(body)
    monthly = body.original_value / (years * 12) if years > 0 else 0
    purchase_date = date.fromisoformat(body.purchase_date) if body.purchase_date else None
    period = await _asset_current_period(db, book_id)
    accumulated = _elapsed_depre_amount(body.original_value, purchase_date, years, period)
    a = FinFixedAsset(
        book_id=book_id, asset_code=body.asset_code or await _next_asset_code(db, book_id), asset_name=body.asset_name,
        category=body.category,
        purchase_date=purchase_date,
        original_value=body.original_value, net_value=max(body.original_value - accumulated, 0),
        depre_method=body.depre_method, depre_years=years,
        monthly_depre=round(monthly, 2), accumulated_depre=accumulated, status=body.status or "in_use",
    )
    db.add(a)
    await db.flush()
    return {"id": a.id, "asset_code": a.asset_code}


@router.put("/assets/{asset_id}")
async def update_asset(
    asset_id: int, body: AssetCreate, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    result = await db.execute(select(FinFixedAsset).where(
        FinFixedAsset.id == asset_id, FinFixedAsset.book_id == book_id
    ))
    asset = result.scalar_one_or_none()
    if not asset:
        raise HTTPException(404, "资产不存在")

    years = _asset_years(body)
    monthly = body.original_value / (years * 12) if years > 0 else 0
    period = await _asset_current_period(db, book_id)
    purchase_date = date.fromisoformat(body.purchase_date) if body.purchase_date else None
    accumulated = _elapsed_depre_amount(body.original_value, purchase_date, years, period)
    asset.asset_code = body.asset_code or asset.asset_code
    asset.asset_name = body.asset_name
    asset.category = body.category
    asset.purchase_date = purchase_date
    asset.original_value = body.original_value
    asset.accumulated_depre = max(float(asset.accumulated_depre or 0), accumulated)
    asset.net_value = max(body.original_value - float(asset.accumulated_depre or 0), 0)
    asset.depre_method = body.depre_method
    asset.depre_years = years
    asset.monthly_depre = round(monthly, 2)
    if body.status:
        asset.status = body.status
    await db.flush()
    return {"ok": True, "id": asset.id, "asset_code": asset.asset_code}


# ── 发票 ──

@router.get("/invoices")
async def list_invoices(
    book_id: int = Query(1),
    invoice_type: Optional[str] = Query(None),
    direction: Optional[str] = Query(None),
    period: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    q = select(FinInvoice).where(FinInvoice.book_id == book_id)
    type_filter = _invoice_type(invoice_type or direction) if (invoice_type or direction) else None
    if type_filter:
        q = q.where(FinInvoice.invoice_type == type_filter)
    if period:
        q = q.where(text("to_char(invoice_date, 'YYYY-MM') = :period")).params(period=period)
    q = q.order_by(FinInvoice.invoice_date.desc(), FinInvoice.id.desc())
    result = await db.execute(q)
    total_amount = total_tax = 0
    rows = []
    for inv in result.scalars():
        total_amount += float(inv.amount or 0)
        total_tax += float(inv.tax_amount or 0)
        direction_value = "in" if inv.invoice_type == "input" else "out"
        rows.append({
            "id": inv.id, "invoice_type": inv.invoice_type, "direction": direction_value,
            "invoice_no": inv.invoice_no,
            "invoice_date": str(inv.invoice_date) if inv.invoice_date else "",
            "counterpart": inv.counterpart, "counterparty": inv.counterpart,
            "amount": float(inv.amount or 0),
            "tax_amount": float(inv.tax_amount or 0), "total_amount": float(inv.total_amount or 0),
            "status": inv.status,
        })
    return {"rows": rows, "total_amount": round(total_amount, 2), "total_tax": round(total_tax, 2)}


def _invoice_type(value: Optional[str]) -> str:
    if value in ("in", "input"):
        return "input"
    if value in ("out", "output"):
        return "output"
    return "output"


class InvoiceCreate(BaseModel):
    invoice_type: Optional[str] = None
    direction: Optional[str] = None
    invoice_no: Optional[str] = None
    invoice_date: Optional[str] = None
    counterpart: Optional[str] = None
    counterparty: Optional[str] = None
    amount: float = 0
    tax_amount: float = 0
    tax_rate: float = 0
    remark: Optional[str] = None

@router.post("/invoices", status_code=201)
async def create_invoice(
    body: InvoiceCreate, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    inv = FinInvoice(
        book_id=book_id, invoice_type=_invoice_type(body.invoice_type or body.direction),
        invoice_no=body.invoice_no,
        invoice_date=date.fromisoformat(body.invoice_date) if body.invoice_date else None,
        counterpart=body.counterpart or body.counterparty, amount=body.amount,
        tax_amount=body.tax_amount, total_amount=body.amount + body.tax_amount,
        tax_rate=body.tax_rate, remark=body.remark,
    )
    db.add(inv)
    await db.flush()
    return {"id": inv.id}


@router.put("/invoices/{invoice_id}/verify")
async def verify_invoice(
    invoice_id: int, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    result = await db.execute(select(FinInvoice).where(FinInvoice.id == invoice_id, FinInvoice.book_id == book_id))
    inv = result.scalar_one_or_none()
    if not inv:
        raise HTTPException(404, "发票不存在")
    inv.status = "verified"
    await db.flush()
    return {"ok": True}


@router.delete("/invoices/{invoice_id}")
async def delete_invoice(
    invoice_id: int, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    result = await db.execute(select(FinInvoice).where(FinInvoice.id == invoice_id, FinInvoice.book_id == book_id))
    inv = result.scalar_one_or_none()
    if not inv:
        raise HTTPException(404, "发票不存在")
    await db.delete(inv)
    await db.flush()
    return {"ok": True}


# ── 工资 ──

@router.get("/payroll")
async def list_payroll(
    book_id: int = Query(1),
    pay_month: Optional[str] = Query(None),
    period: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    pay_month = pay_month or period
    if not pay_month:
        today = date.today()
        pay_month = f"{today.year}-{today.month:02d}"
    result = await db.execute(
        select(FinPayrollRecord)
        .where(FinPayrollRecord.book_id == book_id, FinPayrollRecord.pay_month == pay_month)
        .order_by(FinPayrollRecord.department, FinPayrollRecord.employee_name)
    )
    rows = []
    summary = {"total_gross": 0, "total_net": 0, "total_tax": 0, "count": 0}
    for p in result.scalars():
        summary["total_gross"] += float(p.gross_pay or 0)
        summary["total_net"] += float(p.net_pay or 0)
        summary["total_tax"] += float(p.tax or 0)
        summary["count"] += 1
        deduction = float(p.deductions or 0)
        rows.append({
            "id": p.id, "employee_name": p.employee_name, "department": p.department,
            "base_salary": float(p.base_salary or 0), "bonus": float(p.bonus or 0),
            "deductions": deduction, "deduction": deduction,
            "social_insurance": float(p.social_insurance or 0),
            "housing_fund": float(p.housing_fund or 0),
            "tax": float(p.tax or 0),
            "gross_pay": float(p.gross_pay or 0), "net_pay": float(p.net_pay or 0),
            "status": p.status,
        })
    return {"rows": rows, "summary": summary}


class PayrollCreate(BaseModel):
    employee_name: str
    department: Optional[str] = None
    base_salary: float = 0
    bonus: float = 0
    deductions: Optional[float] = None
    deduction: Optional[float] = None
    social_insurance: float = 0
    housing_fund: float = 0
    tax: float = 0
    net_pay: Optional[float] = None
    status: str = "pending"
    cost_type: str = "管理费用"


def _payroll_amounts(body: PayrollCreate) -> dict:
    deductions = body.deductions if body.deductions is not None else (body.deduction or 0)
    gross = float(body.base_salary or 0) + float(body.bonus or 0)
    calculated_net = gross - float(deductions or 0) - float(body.social_insurance or 0) - float(body.housing_fund or 0) - float(body.tax or 0)
    return {
        "deductions": float(deductions or 0),
        "gross_pay": round(gross, 2),
        "net_pay": round(float(body.net_pay) if body.net_pay is not None else max(calculated_net, 0), 2),
    }


@router.post("/payroll", status_code=201)
async def create_payroll(
    body: PayrollCreate, book_id: int = Query(1),
    pay_month: Optional[str] = Query(None),
    period: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    pay_month = pay_month or period
    if not pay_month:
        today = date.today()
        pay_month = f"{today.year}-{today.month:02d}"
    exists = await db.execute(select(FinPayrollRecord.id).where(
        FinPayrollRecord.book_id == book_id,
        FinPayrollRecord.pay_month == pay_month,
        FinPayrollRecord.employee_name == body.employee_name,
    ))
    if exists.scalar_one_or_none():
        raise HTTPException(400, "该员工本月工资已存在，请编辑原记录")
    amounts = _payroll_amounts(body)
    rec = FinPayrollRecord(
        book_id=book_id, pay_month=pay_month, employee_name=body.employee_name,
        department=body.department, base_salary=body.base_salary, bonus=body.bonus,
        deductions=amounts["deductions"], social_insurance=body.social_insurance,
        housing_fund=body.housing_fund, tax=body.tax,
        gross_pay=amounts["gross_pay"], net_pay=amounts["net_pay"],
        cost_type=body.cost_type, status=body.status,
    )
    db.add(rec)
    await db.flush()
    return {"id": rec.id}


@router.put("/payroll/{payroll_id}")
async def update_payroll(
    payroll_id: int, body: PayrollCreate, book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    result = await db.execute(select(FinPayrollRecord).where(
        FinPayrollRecord.id == payroll_id, FinPayrollRecord.book_id == book_id
    ))
    rec = result.scalar_one_or_none()
    if not rec:
        raise HTTPException(404, "工资记录不存在")
    amounts = _payroll_amounts(body)
    rec.employee_name = body.employee_name
    rec.department = body.department
    rec.base_salary = body.base_salary
    rec.bonus = body.bonus
    rec.deductions = amounts["deductions"]
    rec.social_insurance = body.social_insurance
    rec.housing_fund = body.housing_fund
    rec.tax = body.tax
    rec.gross_pay = amounts["gross_pay"]
    rec.net_pay = amounts["net_pay"]
    rec.cost_type = body.cost_type
    rec.status = body.status
    await db.flush()
    return {"ok": True, "id": rec.id}



# ═══════════════════════════════════════════════════════════
# 工资→凭证
# ═══════════════════════════════════════════════════════════

@router.post("/payroll/generate-voucher")
async def payroll_generate_voucher(
    book_id: int = Query(1),
    pay_month: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """
    工资表→凭证：汇总当月工资，按 cost_type 分科目
    管理费用 → 借 6602管理费用 / 贷 2211应付职工薪酬
    销售费用 → 借 6601销售费用 / 贷 2211应付职工薪酬
    """
    result = await db.execute(text("""
        SELECT cost_type, SUM(gross_pay) AS total
        FROM fin_payroll_records
        WHERE book_id = :bid AND pay_month = :m AND gross_pay > 0
        GROUP BY cost_type
    """), {"bid": book_id, "m": pay_month})
    groups = list(result.mappings())
    if not groups:
        return {"ok": False, "error": "当月无工资数据"}

    # 科目映射
    type_to_code = {"管理费用": "6602", "销售费用": "6601"}
    acct_ids = {}
    for code in ["6602", "6601", "2211"]:
        ar = await db.execute(text("SELECT id FROM fin_accounts WHERE book_id=:bid AND account_code=:c"), {"bid": book_id, "c": code})
        aid = ar.scalar()
        if aid:
            acct_ids[code] = aid

    if "2211" not in acct_ids:
        return {"ok": False, "error": "科目 2211应付职工薪酬 不存在"}

    from app.models.finance.accounting import FinVoucher, FinVoucherLine
    from app.api.v1.finance.vouchers import next_voucher_no
    no_data = await next_voucher_no(book_id, db, current_user)

    total_amount = sum(float(g["total"] or 0) for g in groups)
    v = FinVoucher(
        book_id=book_id, voucher_no=no_data["next_no"], voucher_type="记",
        voucher_date=date.today(), source_type="payroll", source_ref=pay_month,
        summary=f"{pay_month} 工资计提", status="draft",
        created_by=str(current_user.id),
        total_debit=total_amount, total_credit=total_amount,
    )
    db.add(v)
    await db.flush()

    line_no = 1
    for g in groups:
        code = type_to_code.get(g["cost_type"], "6602")
        if code in acct_ids:
            db.add(FinVoucherLine(voucher_id=v.id, line_no=line_no, account_id=acct_ids[code],
                                   account_code=code, debit_amount=float(g["total"]), credit_amount=0,
                                   summary=f"{g['cost_type']}（工资）"))
            line_no += 1

    db.add(FinVoucherLine(voucher_id=v.id, line_no=line_no, account_id=acct_ids["2211"],
                           account_code="2211", debit_amount=0, credit_amount=total_amount,
                           summary="应付职工薪酬"))

    from app.models.finance.settings import FinAuditLog
    db.add(FinAuditLog(book_id=book_id, operator=str(current_user.id),
                        action="payroll_voucher", target=f"voucher:{v.id}", detail=f"{pay_month} 总额={total_amount}"))

    return {"ok": True, "voucher_id": v.id, "voucher_no": v.voucher_no, "total": total_amount}


# ═══════════════════════════════════════════════════════════
# 折旧→凭证
# ═══════════════════════════════════════════════════════════

@router.post("/assets/depreciate")
async def assets_depreciate(
    book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """
    计提本月折旧：遍历在用资产，累加月折旧，生成凭证
    借 6602管理费用（折旧） / 贷 1602累计折旧
    """
    result = await db.execute(
        select(FinFixedAsset).where(
            FinFixedAsset.book_id == book_id,
            FinFixedAsset.status == "in_use",
            FinFixedAsset.monthly_depre > 0,
        )
    )
    assets = list(result.scalars())
    if not assets:
        return {"ok": False, "error": "无可折旧资产"}

    total_depre = sum(float(a.monthly_depre or 0) for a in assets)

    # 更新累计折旧和净值
    for a in assets:
        md = float(a.monthly_depre or 0)
        a.accumulated_depre = (a.accumulated_depre or 0) + md
        a.net_value = float(a.original_value or 0) - float(a.accumulated_depre or 0)
        if a.net_value <= 0:
            a.net_value = 0
            a.status = "scrapped"

    # 获取科目
    acct_ids = {}
    for code in ["6602", "1602"]:
        ar = await db.execute(text("SELECT id FROM fin_accounts WHERE book_id=:bid AND account_code=:c"), {"bid": book_id, "c": code})
        acct_ids[code] = ar.scalar()

    if not acct_ids.get("6602") or not acct_ids.get("1602"):
        return {"ok": False, "error": "科目 6602 或 1602 不存在"}

    from app.models.finance.accounting import FinVoucher, FinVoucherLine
    from app.api.v1.finance.vouchers import next_voucher_no
    no_data = await next_voucher_no(book_id, db, current_user)

    v = FinVoucher(
        book_id=book_id, voucher_no=no_data["next_no"], voucher_type="记",
        voucher_date=date.today(), source_type="depreciation",
        summary=f"本月折旧 {len(assets)}项资产", status="draft",
        created_by=str(current_user.id),
        total_debit=round(total_depre, 2), total_credit=round(total_depre, 2),
    )
    db.add(v)
    await db.flush()

    db.add(FinVoucherLine(voucher_id=v.id, line_no=1, account_id=acct_ids["6602"],
                           account_code="6602", debit_amount=round(total_depre, 2), credit_amount=0,
                           summary=f"折旧费（{len(assets)}项）"))
    db.add(FinVoucherLine(voucher_id=v.id, line_no=2, account_id=acct_ids["1602"],
                           account_code="1602", debit_amount=0, credit_amount=round(total_depre, 2),
                           summary="累计折旧"))

    from app.models.finance.settings import FinAuditLog
    db.add(FinAuditLog(book_id=book_id, operator=str(current_user.id),
                        action="depreciation", target=f"voucher:{v.id}",
                        detail=f"{len(assets)}项资产,总折旧={total_depre:.2f}"))

    return {"ok": True, "voucher_id": v.id, "voucher_no": v.voucher_no,
            "asset_count": len(assets), "total_depre": round(total_depre, 2)}


# ═══════════════════════════════════════════════════════════
# 工资批量导入
# ═══════════════════════════════════════════════════════════

from fastapi import UploadFile, File as FileParam

@router.post("/payroll/import")
async def import_payroll(
    file: UploadFile = FileParam(...),
    book_id: int = Query(1),
    pay_month: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """批量导入工资（Excel: 姓名/部门/基本工资/奖金/扣款/社保/公积金/个税）"""
    from openpyxl import load_workbook
    import io
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    imported = 0
    errors = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or not row[0]:
            continue
        try:
            name = str(row[0]).strip()
            dept = str(row[1]).strip() if row[1] else None
            base = float(row[2] or 0)
            bonus = float(row[3] or 0)
            deductions = float(row[4] or 0)
            social = float(row[5] or 0) if len(row) > 5 else 0
            housing = float(row[6] or 0) if len(row) > 6 else 0
            tax = float(row[7] or 0) if len(row) > 7 else 0
            gross = base + bonus
            net = gross - deductions - social - housing - tax
            db.add(FinPayrollRecord(
                book_id=book_id, pay_month=pay_month, employee_name=name,
                department=dept, base_salary=base, bonus=bonus,
                deductions=deductions, social_insurance=social,
                housing_fund=housing, tax=tax,
                gross_pay=gross, net_pay=max(net, 0),
            ))
            imported += 1
        except Exception as e:
            errors.append(f"第{idx}行: {e}")
    if imported > 0:
        await db.flush()
    return {"imported": imported, "errors": errors[:10]}



# ═══════════════════════════════════════════════════════════
# AI 财务诊断摘要
# ═══════════════════════════════════════════════════════════

@router.post("/ai/diagnosis")
async def ai_finance_diagnosis(
    book_id: int = Query(1),
    month: str = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """
    AI 财务诊断：汇总当月财务数据，调用 LLM 生成经营分析报告。
    结果写入 fin_ai_reports，不自动推送。
    """
    from datetime import date as _date
    if not month:
        today = _date.today()
        month = f"{today.year}-{today.month:02d}"

    # 汇总当月数据
    from app.services import finance_service
    overview = await finance_service.get_overview(db, month)
    ranking = await finance_service.get_store_ranking(db, month, limit=10)

    # 构造 prompt
    top_stores = "\n".join([
        f"  {i+1}. {s['store_name']} 销售¥{s['sales']} 利润¥{s['gross_profit']} 利润率{s['profit_rate']}%"
        for i, s in enumerate(ranking[:5])
    ])
    prompt = f"""你是一位资深的服装行业财务分析师。请根据以下{month}月财务数据，给出经营分析和建议：

【月度财务概况】
- 销售收入: ¥{overview['total_sales']:,.2f}
- 退款金额: ¥{overview['total_refund']:,.2f}
- 毛利: ¥{overview['gross_profit']:,.2f}
- 净利: ¥{overview['net_profit']:,.2f}
- 净利率: {overview['profit_rate']}%
- 广告费: ¥{overview['total_ad_cost']:,.2f}
- 物流费: ¥{overview['total_logistics']:,.2f}
- 活跃店铺数: {overview['store_count']}

【TOP5 店铺利润排行】
{top_stores}

请从以下维度分析：
1. 整体经营健康度评估（用 ★ 评分）
2. 利润率分析与风险点
3. 费用结构优化建议
4. 店铺经营差异化分析
5. 下月经营建议（3条关键行动项）

要求：语言简洁专业，结论明确，每个维度控制在3-5句话。"""

    system = "你是服装电商企业的财务总监助理AI。用中文回答，结构化输出，实事求是。"

    try:
        from app.services.llm_service import llm_service
        content_text = await llm_service.chat(prompt, system=system, temperature=0.4)

        # 清洗思考链
        import re
        content_text = re.sub(r'<think>.*?</think>', '', content_text, flags=re.DOTALL).strip()

        # 保存到 fin_ai_reports
        from app.models.finance.settings import FinAiReport
        report = FinAiReport(
            book_id=book_id, report_type="monthly",
            report_date=month, content=content_text, status="completed",
        )
        db.add(report)
        await db.flush()

        return {"ok": True, "report_id": report.id, "content": content_text}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@router.get("/ai/reports")
async def list_ai_reports(
    book_id: int = Query(1),
    limit: int = Query(10),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """AI 财务报告列表"""
    from app.models.finance.settings import FinAiReport
    result = await db.execute(
        select(FinAiReport).where(FinAiReport.book_id == book_id)
        .order_by(FinAiReport.created_at.desc()).limit(limit)
    )
    return {"rows": [
        {"id": r.id, "report_type": r.report_type, "report_date": r.report_date,
         "content": r.content or "", "status": r.status,
         "created_at": str(r.created_at) if r.created_at else ""}
        for r in result.scalars()
    ]}


# ═══════════════════════════════════════════════════════════
# 企微财务日报推送
# ═══════════════════════════════════════════════════════════

@router.post("/push/daily-report")
async def push_finance_daily(
    month: str = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_user),
):
    """
    推送财务月报到企业微信群（使用 .env 中 WECOM_F_GROUP_WEBHOOK）。
    默认不自动触发，需手动调用。
    """
    from datetime import date as _date
    if not month:
        today = _date.today()
        month = f"{today.year}-{today.month:02d}"

    from app.services import finance_service
    overview = await finance_service.get_overview(db, month)

    markdown_content = f"""## 💰 财务月报 {month}

> **销售收入**: ¥{overview['total_sales']:,.0f}
> **退款金额**: ¥{overview['total_refund']:,.0f}
> **毛利**: ¥{overview['gross_profit']:,.0f}
> **净利**: ¥{overview['net_profit']:,.0f}
> **净利率**: {overview['profit_rate']}%
> **活跃店铺**: {overview['store_count']}家

**费用明细**:
- 广告: ¥{overview['total_ad_cost']:,.0f}
- 物流: ¥{overview['total_logistics']:,.0f}
- 包材: ¥{overview['total_pack_cost']:,.0f}
- 人工: ¥{overview['total_labor_cost']:,.0f}"""

    # 推送
    from app.config.settings import settings
    webhook = settings.WECOM_F_GROUP_WEBHOOK
    if not webhook:
        return {"ok": False, "error": "未配置 WECOM_F_GROUP_WEBHOOK"}

    import httpx
    payload = {"msgtype": "markdown", "markdown": {"content": markdown_content}}
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            resp = await client.post(webhook, json=payload)
            data = resp.json()
            ok = data.get("errcode", -1) == 0

            # 写推送日志
            from app.models.push_log import MsgPushLog
            db.add(MsgPushLog(
                push_type="wecom_webhook", target="finance_group",
                content=markdown_content[:500], status="success" if ok else "failed",
                error_msg=str(data) if not ok else None,
            ))
            return {"ok": ok, "errcode": data.get("errcode")}
    except Exception as e:
        return {"ok": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════
# 银行对账
# ═══════════════════════════════════════════════════════════

@router.get("/bank-reconciliation")
async def bank_reconciliation(
    book_id: int = Query(1),
    account_id: int = Query(None),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """
    银行余额调节表：对比出纳账户余额与总账 1002银行存款 余额。
    """
    if not account_id:
        return {"error": "请选择账户", "rows": []}

    from datetime import date as _date
    today = _date.today()
    period = f"{today.year}-{today.month:02d}"

    # 出纳账户余额
    ca_r = await db.execute(text(
        "SELECT account_name, balance FROM fin_cash_accounts WHERE id = :aid AND book_id = :bid"
    ), {"aid": account_id, "bid": book_id})
    ca = ca_r.mappings().first()
    cashier_balance = float(ca["balance"]) if ca else 0
    account_name = ca["account_name"] if ca else ""

    # 总账 1002 余额
    ledger_r = await db.execute(text("""
        SELECT COALESCE(closing_debit, 0) - COALESCE(closing_credit, 0) AS balance
        FROM fin_ledger_balances lb
        JOIN fin_accounts a ON a.id = lb.account_id
        WHERE lb.book_id = :bid AND a.account_code = '1002' AND lb.period = :p
    """), {"bid": book_id, "p": period})
    ledger_balance = float(ledger_r.scalar() or 0)

    diff = round(cashier_balance - ledger_balance, 2)

    # 未达账项（出纳有但总账没有的流水）
    unmatched_r = await db.execute(text("""
        SELECT flow_date, flow_type, amount, category, counterpart, remark
        FROM fin_cash_flows
        WHERE book_id = :bid AND cash_account_id = :aid AND voucher_id IS NULL
        ORDER BY flow_date DESC LIMIT 50
    """), {"bid": book_id, "aid": account_id})

    return {
        "account_name": account_name,
        "cashier_balance": cashier_balance,
        "ledger_balance": ledger_balance,
        "difference": diff,
        "is_balanced": abs(diff) < 0.01,
        "unmatched_items": [dict(r) for r in unmatched_r.mappings()],
    }


# ═══════════════════════════════════════════════════════════
# 税务配置与台账
# ═══════════════════════════════════════════════════════════

@router.get("/tax/config")
async def list_tax_config(
    book_id: int = Query(1),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """税种配置列表（从应交税费子科目推导）"""
    result = await db.execute(text("""
        SELECT a.id, a.account_code, a.account_name,
               COALESCE(lb.period_credit - lb.period_debit, 0) AS current_amount
        FROM fin_accounts a
        LEFT JOIN fin_ledger_balances lb ON lb.account_id = a.id AND lb.book_id = :bid
        WHERE a.book_id = :bid AND (a.account_code = '2221' OR a.parent_id = (
            SELECT id FROM fin_accounts WHERE book_id = :bid AND account_code = '2221'
        ))
        ORDER BY a.account_code
    """), {"bid": book_id})
    return {"rows": [
        {"id": r["id"], "code": r["account_code"], "name": r["account_name"],
         "current_amount": float(r["current_amount"] or 0)}
        for r in result.mappings()
    ]}


@router.get("/tax/summary")
async def tax_summary(
    book_id: int = Query(1),
    period: str = Query(None),
    db: AsyncSession = Depends(get_db),
    _ = Depends(get_current_user),
):
    """税务汇总：当期应交税费余额 + 各税种明细"""
    from datetime import date as _date
    if not period:
        today = _date.today()
        period = f"{today.year}-{today.month:02d}"

    # 2221 应交税费余额
    tax_r = await db.execute(text("""
        SELECT COALESCE(lb.closing_credit - lb.closing_debit, 0) AS balance
        FROM fin_accounts a
        JOIN fin_ledger_balances lb ON lb.account_id = a.id AND lb.period = :p AND lb.book_id = :bid
        WHERE a.book_id = :bid AND a.account_code = '2221'
    """), {"bid": book_id, "p": period})
    tax_balance = float(tax_r.scalar() or 0)

    # 发票汇总
    inv_r = await db.execute(text("""
        SELECT invoice_type, COUNT(*) AS cnt, COALESCE(SUM(tax_amount), 0) AS total_tax
        FROM fin_invoices WHERE book_id = :bid
        GROUP BY invoice_type
    """), {"bid": book_id})
    invoice_summary = {r["invoice_type"]: {"count": r["cnt"], "tax": float(r["total_tax"])} for r in inv_r.mappings()}

    return {
        "period": period,
        "tax_payable": tax_balance,
        "output_tax": invoice_summary.get("output", {"count": 0, "tax": 0}),
        "input_tax": invoice_summary.get("input", {"count": 0, "tax": 0}),
        "net_tax": invoice_summary.get("output", {"tax": 0})["tax"] - invoice_summary.get("input", {"tax": 0})["tax"],
    }
